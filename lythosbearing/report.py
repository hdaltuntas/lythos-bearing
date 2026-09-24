"""
Calculation report for Lythos Bearing.

The report is assembled as HTML (inputs, geometry and stresses, the capacity
of every method, the factors behind the governing one, the checks, the
figures, the warnings, the method notes and, if one was run, the study) and
exported as
  * PDF   — through reportlab (see `lythosbearing.pdf`)
  * HTML  — a single self-contained file (figures embedded as base64)
  * DOCX  — optional, needs `python-docx`

There is one assembly, `build_html()`, so all three formats say the same thing.
"""

from __future__ import annotations

import base64
import datetime
import html
import io
import math
from typing import Any, Dict, List, Optional

from matplotlib.figure import Figure

from . import study_plots
from .config import APP_NAME, APP_VERSION
from .i18n import TRANSLATIONS, warning_text
from .plotting import PLOT_KEYS, Plotter
from .study import OUTPUTS
from .summary import method_label

TEXTS = {
    "en": {
        "title": f"{APP_NAME} — Calculation Report",
        "date": "Date", "analyst": "Analyst", "software": "Software",
        "sec_inputs": "1. Input data", "sec_found": "1.1 Foundation and actions",
        "sec_soil": "1.2 Soil profile", "sec_opts": "1.3 Method, options and criteria",
        "sec_geom": "2. Geometry and stresses",
        "sec_ecc": "2.1 Eccentricity and effective area",
        "sec_contact": "2.2 Contact pressure",
        "sec_zone": "2.3 Failure zone and design strength",
        "sec_capacity": "3. Bearing capacity", "sec_methods": "3.1 By method",
        "sec_factors": "3.2 Factors of the governing method",
        "sec_terms": "3.3 The three terms", "sec_other": "3.4 Other methods",
        "sec_checks": "4. Checks", "sec_ec7": "4.1 EN 1997-1 verification",
        "sec_width": "4.2 The width the checks need",
        "sec_figs": "5. Figures", "sec_warn": "6. Warnings", "sec_notes": "7. Method notes",
        "parameter": "Parameter", "value": "Value", "unit": "Unit",
        "shape": "Shape", "B": "Width B", "L": "Length L", "D": "Diameter D",
        "Df": "Foundation depth Df", "tilt": "Base tilt η", "slope": "Ground slope β",
        "V": "Vertical load V", "Hb": "Horizontal load along B",
        "Hl": "Horizontal load along L", "Mb": "Moment about the L axis",
        "Ml": "Moment about the B axis", "varfrac": "Variable share of the actions",
        "zw": "Water table depth", "gw": "Unit weight of water γw",
        "yes": "yes", "no": "no",
        "layer": "Layer", "type": "Type", "top": "Top", "bottom": "Bottom",
        "method": "Method", "analysis": "Analysis", "shear": "Failure mode",
        "factors_used": "Corrections applied", "layer_model": "Layer model",
        "zone_factor": "Failure zone factor", "approach": "Verification",
        "FS": "Required factor of safety", "FS_sliding": "Required FS against sliding",
        "ecc_limit": "Allowable eccentricity B/x", "seismic": "Earthquake",
        "kh": "Horizontal coefficient kh", "kv": "Vertical coefficient kv",
        "eB": "Eccentricity e_B", "eL": "Eccentricity e_L",
        "Beff": "Effective width B'", "Leff": "Effective length L'",
        "Aeff": "Effective area A'", "qmax": "Maximum contact pressure",
        "qmin": "Minimum contact pressure", "kern": "Within the middle third",
        "contact": "Length in contact",
        "sigma_v": "Total overburden at the base σv0", "u": "Pore pressure u",
        "sigma_eff": "Effective overburden at the base σ'v0",
        "zone": "Failure zone below the base", "zone_layers": "Layers in the zone",
        "c_design": "Design cohesion c'", "phi_design": "Design friction angle φ'",
        "cu_design": "Design undrained strength cu",
        "gamma_design": "Unit weight below the base",
        "qult": "q_ult", "qnet": "q_net,ult", "qall": "q_all", "applied": "Applied",
        "Nc": "Nc", "Nq": "Nq", "Ng": "Nγ", "factor": "Factor",
        "c_term": "cohesion", "q_term": "surcharge", "g_term": "self weight",
        "term": "Term", "share": "Share",
        "check": "Check", "actual": "Actual", "allowable": "Allowable",
        "status": "Status", "c_bearing": "Bearing capacity", "c_sliding": "Sliding",
        "c_ecc": "Eccentricity", "combination": "Combination", "sets": "Factor sets",
        "Ed": "Ed", "Rd": "Rd", "util": "Λ = Ed/Rd",
        "ok": "OK", "notok": "NOT OK", "na": "n/a",
        "sec_study": "8. Parametric / reliability study", "sec_study_vars": "8.1 Variables",
        "sec_study_stats": "8.2 Statistics of the outputs",
        "sec_study_rel": "8.3 Probability of failure",
        "sec_study_sens": "8.4 Sensitivities (Spearman ρ)", "sec_study_figs": "8.5 Figures",
        "study_method": "Sampling method", "study_n": "Samples",
        "study_ok": "Successful analyses", "variable": "Variable", "mode": "Mode",
        "range": "Range", "dist": "Distribution", "mean": "Mean", "cov": "CoV",
        "min": "Min", "max": "Max", "output": "Output", "std": "Std",
        "criterion": "Criterion", "n_fail": "Failed", "pf": "P", "pf_ci": "95 % CI",
        "beta_idx": "β",
        "notes": [
            "The capacity is the general bearing capacity equation "
            "q_ult = c·Nc·sc·dc·ic·bc·gc + q·Nq·sq·dq·iq·bq·gq + ½·γ·B'·Nγ·sγ·dγ·iγ·bγ·gγ, "
            "each method bringing its own bearing capacity factors and its own "
            "corrections. A correction the method does not define is reported as 1.",
            "Terzaghi's method carries shape factors only; Meyerhof's adds depth and load "
            "inclination; Hansen's and Vesić's add base tilt and ground slope; "
            "EN 1997-1 Annex D gives shape, load inclination and base inclination factors "
            "and no depth factors.",
            "At φ' = 0 Hansen's expression is additive, q_ult = 5.14·cu·(1 + s'c + d'c "
            "− i'c − b'c − g'c) + q, and is written out that way rather than multiplied.",
            "The surcharge q is the effective overburden at the founding depth in a "
            "drained analysis and the total overburden in an undrained one; the unit "
            "weight below the base is the buoyant weight where the soil is submerged.",
            "The strength is averaged over the failure zone, whose depth is Prandtl's "
            "(B'/2)·cos φ/cos(45 + φ/2)·e^((π/4 + φ/2)·tan φ) — 0.707·B' at φ' = 0 — the "
            "friction angle through its tangent and everything else directly. The zone "
            "and the strength that sets it are settled by repetition.",
            "An eccentric load is carried on Meyerhof's effective area, the rectangle "
            "B' = B − 2·e_B by L' = L − 2·e_L centred on the resultant; for a circle it is "
            "the equivalent rectangle of the loaded segment. The contact pressure over the "
            "whole base is reported beside it, and is triangular once the resultant leaves "
            "the middle third.",
            "The factor of safety is taken on the net capacity, FS = q_net,ult / q_net; "
            "the allowable pressure is q_net,ult/FS + q. A Design Approach instead factors "
            "the actions, the strength and the resistance with the recommended values of "
            "EN 1997-1 Annex A, and compares Ed with Rd.",
            "Sliding is checked as R = V·tan δ + A'·c'a drained, R = A'·cu undrained, with "
            "δ = δ/φ'·φ'.",
            "The required width is the smallest width that satisfies the bearing check, "
            "found by bisection with every other input held as it is; a rectangle keeps "
            "its length-to-width ratio.",
        ],
        "note_two_layer": "Two layers: a strong layer of limited thickness over a weak one "
                          "is checked by Meyerhof & Hanna's punching model, q_ult = q_b + "
                          "(1 + B/L)·[2·ca·H/B + γ₁·H²·(1 + 2·Df/H)·Ks·tan φ₁/B] − γ₁·H, "
                          "capped by the capacity of the top layer alone; or by spreading "
                          "the load onto the weak layer. Ks is the value entered, or "
                          "1 − sin φ₁.",
        "note_seismic": "Earthquake: the structure's inertia kh·V is added to the "
                        "horizontal load and the vertical load becomes V·(1 − kv); the "
                        "soil's own inertia reduces the terms by Paolucci & Pecker's "
                        "(1 − kh/tan φ)^0.35 and (1 − 0.32·kh).",
        "note_insitu": "The SPT and CPT rules give the pressure that keeps the settlement "
                       "within the tolerable value, not a rupture capacity; Ménard's rule "
                       "gives an ultimate capacity, q_net,ult = kp·(pl − p0).",
        "note_rock": "Rock: the Hoek–Brown mass is turned into an equivalent c' and φ' "
                     "(Hoek, Carranza-Torres & Corkum 2002) up to σ3 = σci/4 and run "
                     "through the same equation; the CFEM Ksp method gives an allowable "
                     "pressure that already carries a factor of safety of about 3.",
    },
    "tr": {
        "title": f"{APP_NAME} — Hesap Raporu",
        "date": "Tarih", "analyst": "Hazırlayan", "software": "Yazılım",
        "sec_inputs": "1. Girdi verileri", "sec_found": "1.1 Temel ve yükler",
        "sec_soil": "1.2 Zemin profili", "sec_opts": "1.3 Yöntem, seçenekler ve ölçütler",
        "sec_geom": "2. Geometri ve gerilmeler",
        "sec_ecc": "2.1 Dışmerkezlik ve etkin alan",
        "sec_contact": "2.2 Taban basıncı",
        "sec_zone": "2.3 Göçme bölgesi ve hesap dayanımı",
        "sec_capacity": "3. Taşıma gücü", "sec_methods": "3.1 Yönteme göre",
        "sec_factors": "3.2 Belirleyici yöntemin katsayıları",
        "sec_terms": "3.3 Üç terim", "sec_other": "3.4 Diğer yöntemler",
        "sec_checks": "4. Kontroller", "sec_ec7": "4.1 EN 1997-1 kontrolü",
        "sec_width": "4.2 Kontrollerin gerektirdiği genişlik",
        "sec_figs": "5. Şekiller", "sec_warn": "6. Uyarılar", "sec_notes": "7. Yöntem notları",
        "parameter": "Parametre", "value": "Değer", "unit": "Birim",
        "shape": "Şekil", "B": "Genişlik B", "L": "Boy L", "D": "Çap D",
        "Df": "Temel derinliği Df", "tilt": "Taban eğimi η", "slope": "Arazi eğimi β",
        "V": "Düşey yük V", "Hb": "B yönünde yatay yük", "Hl": "L yönünde yatay yük",
        "Mb": "L ekseni etrafındaki moment", "Ml": "B ekseni etrafındaki moment",
        "varfrac": "Yüklerin hareketli payı",
        "zw": "Su tablası derinliği", "gw": "Suyun birim hacim ağırlığı γw",
        "yes": "evet", "no": "hayır",
        "layer": "Tabaka", "type": "Tür", "top": "Üst", "bottom": "Alt",
        "method": "Yöntem", "analysis": "Analiz", "shear": "Göçme biçimi",
        "factors_used": "Uygulanan düzeltmeler", "layer_model": "Tabaka modeli",
        "zone_factor": "Göçme bölgesi katsayısı", "approach": "Kontrol biçimi",
        "FS": "Gerekli güvenlik sayısı", "FS_sliding": "Kaymaya karşı gerekli GS",
        "ecc_limit": "İzin verilen dışmerkezlik B/x", "seismic": "Deprem",
        "kh": "Yatay katsayı kh", "kv": "Düşey katsayı kv",
        "eB": "Dışmerkezlik e_B", "eL": "Dışmerkezlik e_L",
        "Beff": "Etkin genişlik B'", "Leff": "Etkin boy L'", "Aeff": "Etkin alan A'",
        "qmax": "En büyük taban basıncı", "qmin": "En küçük taban basıncı",
        "kern": "Orta üçte bir içinde", "contact": "Temas uzunluğu",
        "sigma_v": "Tabanda toplam örtü yükü σv0", "u": "Boşluk suyu basıncı u",
        "sigma_eff": "Tabanda efektif örtü yükü σ'v0",
        "zone": "Taban altındaki göçme bölgesi", "zone_layers": "Bölgedeki tabakalar",
        "c_design": "Hesap kohezyonu c'", "phi_design": "Hesap sürtünme açısı φ'",
        "cu_design": "Hesap drenajsız dayanımı cu",
        "gamma_design": "Taban altı birim hacim ağırlığı",
        "qult": "q_ult", "qnet": "q_net,ult", "qall": "q_all", "applied": "Uygulanan",
        "Nc": "Nc", "Nq": "Nq", "Ng": "Nγ", "factor": "Katsayı",
        "c_term": "kohezyon", "q_term": "sürşarj", "g_term": "zati ağırlık",
        "term": "Terim", "share": "Pay",
        "check": "Kontrol", "actual": "Oluşan", "allowable": "İzin verilen",
        "status": "Durum", "c_bearing": "Taşıma gücü", "c_sliding": "Kayma",
        "c_ecc": "Dışmerkezlik", "combination": "Kombinasyon", "sets": "Katsayı takımları",
        "Ed": "Ed", "Rd": "Rd", "util": "Λ = Ed/Rd",
        "ok": "UYGUN", "notok": "UYGUN DEĞİL", "na": "—",
        "sec_study": "8. Parametrik / güvenilirlik çalışması",
        "sec_study_vars": "8.1 Değişkenler", "sec_study_stats": "8.2 Çıktıların istatistikleri",
        "sec_study_rel": "8.3 Göçme olasılığı",
        "sec_study_sens": "8.4 Duyarlılıklar (Spearman ρ)", "sec_study_figs": "8.5 Şekiller",
        "study_method": "Örnekleme yöntemi", "study_n": "Örnek sayısı",
        "study_ok": "Başarılı analiz", "variable": "Değişken", "mode": "Mod",
        "range": "Aralık", "dist": "Dağılım", "mean": "Ortalama", "cov": "CoV",
        "min": "Min", "max": "Maks", "output": "Çıktı", "std": "Std",
        "criterion": "Ölçüt", "n_fail": "Göçen", "pf": "P", "pf_ci": "%95 GA",
        "beta_idx": "β",
        "notes": [
            "Taşıma gücü, genel taşıma gücü denklemidir: "
            "q_ult = c·Nc·sc·dc·ic·bc·gc + q·Nq·sq·dq·iq·bq·gq + ½·γ·B'·Nγ·sγ·dγ·iγ·bγ·gγ; "
            "her yöntem kendi taşıma gücü katsayılarını ve kendi düzeltmelerini getirir. "
            "Yöntemin tanımlamadığı düzeltme raporda 1 olarak görünür.",
            "Terzaghi yönteminde yalnızca şekil katsayıları vardır; Meyerhof derinlik ve yük "
            "eğimini ekler; Hansen ve Vesić taban eğimi ile arazi eğimini de ekler; "
            "EN 1997-1 Ek D şekil, yük eğimi ve taban eğimi katsayıları verir, derinlik "
            "katsayısı vermez.",
            "φ' = 0 durumunda Hansen ifadesi toplamsaldır: q_ult = 5.14·cu·(1 + s'c + d'c "
            "− i'c − b'c − g'c) + q; çarpımsal değil, bu biçimde yazılır.",
            "Sürşarj q, drenajlı analizde temel derinliğindeki efektif örtü yükü, drenajsız "
            "analizde toplam örtü yüküdür; taban altındaki birim hacim ağırlığı, su altında "
            "kalan yerde kaldırma kuvveti düşülmüş ağırlıktır.",
            "Dayanım, derinliği Prandtl'ın (B'/2)·cos φ/cos(45 + φ/2)·e^((π/4 + φ/2)·tan φ) "
            "ifadesiyle verilen göçme bölgesinde ortalanır — φ' = 0 için 0.707·B'. Sürtünme "
            "açısı tanjantı üzerinden, diğerleri doğrudan ortalanır. Bölge ile onu belirleyen "
            "dayanım yinelemeyle uzlaştırılır.",
            "Dışmerkez yük, Meyerhof'un etkin alanında taşınır: bileşkeye göre ortalanmış "
            "B' = B − 2·e_B ve L' = L − 2·e_L dikdörtgeni; dairede yüklü parçanın eşdeğer "
            "dikdörtgeni. Tüm taban üzerindeki basınç dağılımı ayrıca verilir ve bileşke orta "
            "üçte birin dışına çıktığında üçgen olur.",
            "Güvenlik sayısı net taşıma gücü üzerinden alınır: FS = q_net,ult / q_net; izin "
            "verilebilir basınç q_net,ult/FS + q'dur. Tasarım Yaklaşımı ise yükleri, dayanımı "
            "ve direnci EN 1997-1 Ek A'nın önerilen değerleriyle çarpar ve Ed ile Rd'yi "
            "karşılaştırır.",
            "Kayma, drenajlı durumda R = V·tan δ + A'·c'a, drenajsız durumda R = A'·cu ile "
            "kontrol edilir; δ = δ/φ'·φ'.",
            "Gerekli genişlik, diğer bütün girdiler sabit tutularak taşıma gücü kontrolünü "
            "sağlayan en küçük genişliktir ve ikiye bölme ile bulunur; dikdörtgen temel "
            "boy/genişlik oranını korur.",
        ],
        "note_two_layer": "İki tabaka: zayıf tabaka üzerindeki sınırlı kalınlıktaki sağlam "
                          "tabaka, Meyerhof & Hanna zımbalama modeliyle kontrol edilir: "
                          "q_ult = q_b + (1 + B/L)·[2·ca·H/B + γ₁·H²·(1 + 2·Df/H)·Ks·tan φ₁/B] "
                          "− γ₁·H; üst tabakanın tek başına taşıma gücü ile sınırlıdır. "
                          "Alternatif olarak yük zayıf tabakaya yayılır. Ks girilen değerdir "
                          "ya da 1 − sin φ₁.",
        "note_seismic": "Deprem: yapının ataleti kh·V yatay yüke eklenir, düşey yük "
                        "V·(1 − kv) olur; zeminin kendi ataleti terimleri Paolucci & Pecker'in "
                        "(1 − kh/tan φ)^0.35 ve (1 − 0.32·kh) katsayılarıyla azaltır.",
        "note_insitu": "SPT ve CPT kuralları, oturmayı izin verilen değerin altında tutan "
                       "basıncı verir; göçme taşıma gücü değildir. Ménard kuralı nihai taşıma "
                       "gücü verir: q_net,ult = kp·(pl − p0).",
        "note_rock": "Kaya: Hoek–Brown kütlesi σ3 = σci/4'e kadar eşdeğer c' ve φ'ye çevrilir "
                     "(Hoek, Carranza-Torres & Corkum 2002) ve aynı denklemde çalıştırılır; "
                     "CFEM Ksp yöntemi zaten yaklaşık 3 güvenlik sayısı taşıyan izin "
                     "verilebilir basınç verir.",
    },
}

#: The interface's palette: warm paper, ink, terracotta; serif headings
_CSS = """
body { font-family: system-ui, -apple-system, 'Segoe UI', 'Helvetica Neue', Arial, sans-serif;
       font-size: 10pt; color: #141413; }
h1, h2, h3 { font-family: 'Tiempos Text', 'Source Serif 4', 'Iowan Old Style', Palatino,
             Georgia, 'DejaVu Serif', serif; font-weight: 500; }
h1 { font-size: 20pt; color: #141413; margin-bottom: 2px; }
h2 { font-size: 14pt; color: #C6613F; border-bottom: 1px solid #E3E0D5; padding-bottom: 2px;
     margin-top: 20px; }
h3 { font-size: 11.5pt; color: #141413; margin-top: 12px; }
table { border-collapse: collapse; margin: 4px 0 8px 0; }
th { background: #F0EEE6; text-align: left; padding: 3px 6px; border: 1px solid #E3E0D5;
     font-size: 9pt; }
td { padding: 3px 6px; border: 1px solid #E3E0D5; font-size: 9pt; }
.ok { color: #3F7F4F; font-weight: bold; } .bad { color: #B0413E; font-weight: bold; }
.meta { color: #73726C; } .note { color: #73726C; font-size: 9pt; }
"""


def _esc(x) -> str:
    return html.escape(str(x))


def _f(x, nd=2) -> str:
    if x is None or (isinstance(x, float) and not math.isfinite(x)):
        return "—"
    return f"{x:,.{nd}f}"


def _status(T, s: str) -> str:
    if s == "N/A":
        return T["na"]
    return (f'<span class="ok">{T["ok"]}</span>' if s == "OK"
            else f'<span class="bad">{T["notok"]}</span>')


def _table(headers: List[str], rows: List[List[Any]], widths: Optional[List[int]] = None) -> str:
    out = ["<table width='100%'>"]
    if widths:
        cells = "".join(f"<th width='{w}%'>{_esc(h)}</th>" for h, w in zip(headers, widths))
    else:
        cells = "".join(f"<th>{_esc(h)}</th>" for h in headers)
    out.append("<tr>" + cells + "</tr>")
    for r in rows:
        out.append("<tr>" + "".join(f"<td>{c}</td>" for c in r) + "</tr>")
    out.append("</table>")
    return "\n".join(out)


def _kv_table(T, rows: List[List[Any]]) -> str:
    return _table([T["parameter"], T["value"], T["unit"]], rows, [50, 36, 14])


# ----------------------------------------------------------------------
def _png(fig: Figure, dpi: int) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi, bbox_inches="tight")
    return buf.getvalue()


def render_figures(analysis, lang: str, keys=PLOT_KEYS, dpi: int = 130) -> Dict[str, bytes]:
    """The report figures as PNG bytes (off-screen, white, untitled)."""
    plotter = Plotter(analysis, lang, "paper", titles=False)
    out = {}
    for key in keys:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        plotter.draw(key, fig)
        out[key] = _png(fig, dpi)
    return out


def render_study_figures(study, lang: str, dpi: int = 110) -> Dict[str, bytes]:
    """Study figures keyed 'study_<view>' (only the views that apply)."""
    if study is None or not study.rows:
        return {}
    L = TRANSLATIONS[lang]
    views = ["oat", "hist"] if study.method == "oat" else ["hist", "scatter", "tornado"]
    draw = {"oat": study_plots.plot_oat, "hist": study_plots.plot_hist,
            "scatter": study_plots.plot_scatter, "tornado": study_plots.plot_tornado}
    out = {}
    for view in views:
        fig = Figure(figsize=(10, 7), dpi=dpi)
        draw[view](fig, study, L, "FS", theme="paper")
        out[f"study_{view}"] = _png(fig, dpi)
    return out


def _study_section(T, L, study, figures, img_src) -> List[str]:
    parts = [f"<h2 style='page-break-before:always'>{T['sec_study']}</h2>"]
    s = study.summary
    parts.append(_kv_table(T, [
        [T["study_method"], L[f"sampling_{study.method}"], ""],
        [T["study_n"], str(s.get("n_total", 0)), ""],
        [T["study_ok"], str(s.get("n_ok", 0)), ""],
    ]))
    parts.append(f"<h3>{T['sec_study_vars']}</h3>")
    rows = []
    for v in study.variables:
        if v.mode == "range":
            rows.append([_esc(v.label), T["range"], _f(v.vmin, 3), _f(v.vmax, 3), "—", "—",
                         "—", v.n_points])
        else:
            rows.append([_esc(v.label), T["dist"], "—", "—", L[f"dist_{v.dist}"],
                         _f(v.mean, 3), _f(v.cov, 3), "—"])
    parts.append(_table([T["variable"], T["mode"], T["min"], T["max"], T["dist"], T["mean"],
                         T["cov"], "n"], rows, [28, 12, 10, 10, 14, 10, 8, 8]))
    parts.append(f"<h3>{T['sec_study_stats']}</h3>")
    rows = [[_esc(L.get(f"out_{k}", k)), st["n"], _f(st["mean"], 3), _f(st["std"], 3),
             _f(st["p5"], 3), _f(st["p50"], 3), _f(st["p95"], 3)]
            for k, _ in OUTPUTS if k in s.get("stats", {}) for st in [s["stats"][k]]]
    parts.append(_table([T["output"], "n", T["mean"], T["std"], "P5", "P50", "P95"],
                        rows, [34, 8, 12, 12, 11, 11, 12]))
    if s.get("reliability"):
        parts.append(f"<h3>{T['sec_study_rel']}</h3>")
        rows = []
        for name, r in s["reliability"].items():
            rows.append([_esc(L.get(name, name)), r["n"], r["n_fail"], f"{r['pf']:.3g}",
                         f"[{r['pf_lo']:.2g}, {r['pf_hi']:.2g}]",
                         _esc(study_plots._beta_text(r["beta"], r["n"]))])
        parts.append(_table([T["criterion"], "n", T["n_fail"], T["pf"], T["pf_ci"],
                             T["beta_idx"]], rows, [30, 10, 12, 14, 20, 14]))
    if s.get("spearman"):
        parts.append(f"<h3>{T['sec_study_sens']}</h3>")
        outs = [k for k, _ in OUTPUTS if k in s["spearman"]]
        rows = [[_esc(v.label)] + [f"{s['spearman'][k].get(v.path, float('nan')):+.2f}"
                                   for k in outs] for v in study.variables]
        parts.append(_table([T["variable"]] + [L.get(f"out_{k}", k) for k in outs], rows))
    parts.append(f"<h3>{T['sec_study_figs']}</h3>")
    for key in figures:
        if key.startswith("study_"):
            parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")
    return parts


def build_html(analysis, lang: str, figures: Dict[str, bytes], img_src=None,
               study=None) -> str:
    """
    Builds the report HTML. img_src(key) -> value for the <img src> attribute;
    the default embeds base64 data URIs (the self-contained HTML file). The PDF
    and DOCX writers pass a `fig://<key>` mapper and resolve the keys against
    the figure dictionary themselves.
    """
    T = TEXTS.get(lang, TEXTS["en"])
    L = TRANSLATIONS.get(lang, TRANSLATIONS["en"])
    a, res = analysis, analysis.results
    info = a.config.get("project_info", {})
    if img_src is None:
        def img_src(key):
            return "data:image/png;base64," + base64.b64encode(figures[key]).decode("ascii")

    geom, pressure, applied = res["geometry"], res["pressure"], res["applied"]
    checks = res["checks"]

    parts = [f"<html><head><meta charset='utf-8'><style>{_CSS}</style></head><body>"]
    parts.append(f"<h1>{_esc(info.get('title') or T['title'])}</h1>")
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    meta = f"{T['date']}: {stamp}"
    if info.get("analyst"):
        meta += f" &nbsp;·&nbsp; {T['analyst']}: {_esc(info['analyst'])}"
    meta += f" &nbsp;·&nbsp; {T['software']}: {APP_NAME} v{APP_VERSION}"
    parts.append(f"<p class='meta'>{meta}</p>")

    # ---------------------------------------------------------------- inputs
    parts.append(f"<h2>{T['sec_inputs']}</h2>")
    parts.append(f"<h3>{T['sec_found']}</h3>")
    rows = [[T["shape"], L["shape_" + a.shape], ""],
            [T["D"] if a.shape == "circle" else T["B"], _f(a.B), "m"]]
    if a.shape == "rectangle":
        rows.append([T["L"], _f(a.L), "m"])
    rows += [[T["Df"], _f(a.Df), "m"]]
    if a.eta:
        rows.append([T["tilt"], _f(a.eta, 1), "°"])
    if a.beta:
        rows.append([T["slope"], _f(a.beta, 1), "°"])
    rows += [[T["V"], _f(a.V, 1), "kN"], [T["Hb"], _f(a.Hb, 1), "kN"]]
    if a.Hl:
        rows.append([T["Hl"], _f(a.Hl, 1), "kN"])
    rows.append([T["Mb"], _f(a.Mb, 1), "kN·m"])
    if a.Ml:
        rows.append([T["Ml"], _f(a.Ml, 1), "kN·m"])
    rows += [[T["varfrac"], _f(a.variable_fraction), ""],
             [T["zw"], _f(a.profile.zw), "m"], [T["gw"], _f(a.profile.gw), "kN/m³"]]
    if a.seismic_on:
        rows += [[T["seismic"], T["yes"], ""], [T["kh"], _f(a.kh, 3), ""],
                 [T["kv"], _f(a.kv, 3), ""]]
    parts.append(_kv_table(T, rows))

    parts.append(f"<h3>{T['sec_soil']}</h3>")
    rows = [[_esc(row["name"]), L["behaviour_" + row["behaviour"]],
             _f(row["top"]), _f(row["bottom"]), _f(row["gamma"], 1),
             _f(row["gamma_sat"], 1), _f(row["c"], 1), _f(row["phi"], 1),
             _f(row["cu"], 1), _f(row["sigma_eff"], 1)]
            for row in res["layers"]]
    parts.append(_table([T["layer"], T["type"], T["top"], T["bottom"], L["col_gamma"],
                         L["col_gamma_sat"], L["col_c"], L["col_phi"], L["col_cu"],
                         "σ'v0"], rows, [18, 10, 7, 7, 10, 10, 9, 8, 9, 12]))

    parts.append(f"<h3>{T['sec_opts']}</h3>")
    used = [L[key] for name, key in
            (("shape", "shape_factors_label"), ("depth", "depth_factors_label"),
             ("inclination", "inclination_factors_label"), ("base", "base_factors_label"),
             ("ground", "ground_factors_label"),
             ("compressibility", "compressibility_label"))
            if getattr(a, f"use_{name}")]
    parts.append(_kv_table(T, [
        [T["method"], method_label(L, a.method), ""],
        [T["analysis"], L["analysis_" + a.analysis], ""],
        [T["shear"], L["shear_" + a.shear], ""],
        [T["factors_used"], ", ".join(used) or "—", ""],
        [T["layer_model"], L["layer_model_" + a.layer_model], ""],
        [T["zone_factor"], _f(a.zone_factor), ""],
        [T["approach"], L["approach_" + a.approach], ""],
        [T["FS"], _f(a.FS), ""], [T["FS_sliding"], _f(a.FS_sliding), ""],
        [T["ecc_limit"], _f(a.ecc_limit, 1), ""],
    ]))

    # ---------------------------------------------------------- geometry
    parts.append(f"<h2>{T['sec_geom']}</h2>")
    parts.append(f"<h3>{T['sec_ecc']}</h3>")
    parts.append(_kv_table(T, [
        [T["eB"], _f(geom["e_B"], 3), "m"], [T["eL"], _f(geom["e_L"], 3), "m"],
        [T["Beff"], _f(geom["B_eff"], 3), "m"], [T["Leff"], _f(geom["L_eff"], 3), "m"],
        [T["Aeff"], _f(geom["A_eff"]), "m²"],
    ]))
    parts.append(f"<h3>{T['sec_contact']}</h3>")
    parts.append(_kv_table(T, [
        [T["qmax"], _f(pressure["q_max"], 1), "kPa"],
        [T["qmin"], _f(pressure["q_min"], 1), "kPa"],
        [T["kern"], T["yes"] if pressure["in_kern"] else T["no"], ""],
        [T["contact"], _f(pressure["contact"]), "m"],
    ]))
    parts.append(f"<h3>{T['sec_zone']}</h3>")
    params = res["params"]["drained"]
    parts.append(_kv_table(T, [
        [T["sigma_v"], _f(res["surcharge"]["total"], 1), "kPa"],
        [T["u"], _f(res["surcharge"]["pore"], 1), "kPa"],
        [T["sigma_eff"], _f(res["surcharge"]["effective"], 1), "kPa"],
        [T["zone"], _f(res["zone"]["depth"]), "m"],
        [T["zone_layers"], _esc(", ".join(params["layers"])), ""],
        [T["c_design"], _f(params["c"], 1), "kPa"],
        [T["phi_design"], _f(params["phi"], 1), "°"],
        [T["cu_design"], _f(res["params"]["undrained"]["cu"], 1), "kPa"],
        [T["gamma_design"], _f(params["gamma"]), "kN/m³"],
    ]))

    # ---------------------------------------------------------- capacity
    parts.append(f"<h2>{T['sec_capacity']}</h2>")
    parts.append(f"<h3>{T['sec_methods']}</h3>")
    rows = []
    for key, entry in res["methods"].items():
        mark = " *" if key == res["primary"] else ""
        rows.append([_esc(method_label(L, key) + mark),
                     L[f"analysis_label_short_{entry['governing']}"],
                     _f(entry["q_ult"], 1), _f(entry["q_net_ult"], 1),
                     _f(entry["N"]["Nc"], 2), _f(entry["N"]["Nq"], 2),
                     _f(entry["N"]["Ngamma"], 2)])
    parts.append(_table([T["method"], T["analysis"], T["qult"], T["qnet"], T["Nc"],
                         T["Nq"], T["Ng"]], rows, [28, 14, 13, 13, 11, 10, 11]))

    primary = res["methods"].get(res["primary"])
    if primary:
        parts.append(f"<h3>{T['sec_factors']}</h3>")
        rows = [[L[f"factor_{name}"], _f(primary["factors"][name]["c"], 3),
                 _f(primary["factors"][name]["q"], 3), _f(primary["factors"][name]["g"], 3)]
                for name in ("shape", "depth", "inclination", "base", "ground",
                             "compressibility")]
        parts.append(_table([T["factor"], T["c_term"], T["q_term"], T["g_term"]], rows,
                            [40, 20, 20, 20]))
        parts.append(f"<h3>{T['sec_terms']}</h3>")
        total = primary["q_ult"] or 1.0
        rows = [[T[f"{part}_term"], _f(primary["terms"][part], 1),
                 f"{100.0 * primary['terms'][part] / total:.0f} %"]
                for part in ("c", "q", "g")]
        rows.append([f"<b>{T['qult']}</b>", f"<b>{_f(primary['q_ult'], 1)}</b>", "100 %"])
        parts.append(_table([T["term"], "kPa", T["share"]], rows, [50, 25, 25]))

    if res["others"]:
        parts.append(f"<h3>{T['sec_other']}</h3>")
        rows = [[_esc(method_label(L, key)), _f(entry.get("q_ult"), 1),
                 _f(entry.get("q_net_ult"), 1), _f(entry.get("q_all"), 1)]
                for key, entry in res["others"].items()]
        parts.append(_table([T["method"], T["qult"], T["qnet"], T["qall"]], rows,
                            [40, 20, 20, 20]))

    # ---------------------------------------------------------- checks
    parts.append(f"<h2>{T['sec_checks']}</h2>")
    bearing, slide, ecc = checks["bearing"], checks["sliding"], checks["eccentricity"]
    rows = [[T["c_bearing"], _f(bearing["actual"]), _f(bearing["allowable"]),
             _status(T, bearing["status"])],
            [T["c_sliding"], _f(slide["actual"]), _f(slide["allowable"]),
             _status(T, slide["status"])],
            [T["c_ecc"], _f(ecc["actual"], 3), _f(ecc["allowable"], 3),
             _status(T, ecc["status"])]]
    parts.append(_table([T["check"], T["actual"], T["allowable"], T["status"]], rows,
                        [40, 20, 20, 20]))
    parts.append(f"<p>{T['qall']} = {_f(bearing['q_all'], 1)} kPa &nbsp;·&nbsp; "
                 f"{T['applied']} = {_f(applied['q'], 1)} kPa</p>")

    ec7 = res.get("eurocode")
    if ec7 and ec7.get("combinations"):
        parts.append(f"<h3>{T['sec_ec7']} — {L['approach_' + ec7['approach']]}</h3>")
        rows = [[_esc(c["name"]), " + ".join(c["sets"]),
                 L[f"analysis_label_short_{c['analysis']}"], _f(c["Ed"], 0),
                 _f(c["Rd"], 0), _f(c["utilisation"], 3), _status(T, c["status"])]
                for c in ec7["combinations"]]
        parts.append(_table([T["combination"], T["sets"], T["analysis"], T["Ed"], T["Rd"],
                             T["util"], T["status"]], rows, [16, 18, 14, 13, 13, 13, 13]))

    parts.append(f"<h3>{T['sec_width']}</h3>")
    width = res["required_width"]
    parts.append(f"<p>{_esc(L['res_required_width'].format(B=width)) if math.isfinite(width) else _esc(L['res_no_width'].format(hi=60.0))}</p>")

    # ---------------------------------------------------------- figures
    parts.append(f"<h2 style='page-break-before:always'>{T['sec_figs']}</h2>")
    for key in PLOT_KEYS:
        if key not in figures:
            continue
        parts.append(f"<p class='lead'>{_esc(L.get('fig_' + key, key))}</p>")
        parts.append(f"<p><img src='{img_src(key)}' width='640'></p>")

    if res["warnings"]:
        parts.append(f"<h2>{T['sec_warn']}</h2><ul>")
        parts.extend(f"<li>{_esc(warning_text(lang, w))}</li>" for w in res["warnings"])
        parts.append("</ul>")

    parts.append(f"<h2>{T['sec_notes']}</h2><ul class='note'>")
    notes = list(T["notes"])
    if a.layer_model == "two_layer":
        notes.append(T["note_two_layer"])
    if a.seismic_on:
        notes.append(T["note_seismic"])
    if a.insitu_on:
        notes.append(T["note_insitu"])
    if a.rock_on:
        notes.append(T["note_rock"])
    parts.extend(f"<li>{_esc(n)}</li>" for n in notes)
    parts.append("</ul>")

    if study is not None and study.rows:
        parts.extend(_study_section(T, L, study, figures, img_src))
    parts.append("</body></html>")
    return "\n".join(parts)


# ----------------------------------------------------------------------
# exporters
# ----------------------------------------------------------------------
def _all_figures(analysis, lang, study):
    figures = render_figures(analysis, lang)
    figures.update(render_study_figures(study, lang))
    return figures


def export_html(path: str, analysis, lang: str, study=None) -> None:
    figures = _all_figures(analysis, lang, study)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write(build_html(analysis, lang, figures, study=study))


def export_pdf(path: str, analysis, lang: str, study=None) -> None:
    """PDF through reportlab; the figures travel as `fig://<key>` references."""
    from . import pdf as pdf_writer

    T = TEXTS.get(lang, TEXTS["en"])
    info = analysis.config.get("project_info", {})
    figures = _all_figures(analysis, lang, study)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}",
                           study=study)
    pdf_writer.html_to_pdf(path, html_text, figures,
                           title=info.get("title", T["title"]),
                           author=info.get("analyst", ""),
                           footer=f"{info.get('title', '')} — {APP_NAME} v{APP_VERSION}")


def export_docx(path: str, analysis, lang: str, study=None) -> None:
    """Word report; requires python-docx."""
    try:
        import docx
        from docx.shared import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-docx is not installed (pip install python-docx).") from exc
    import re

    figures = _all_figures(analysis, lang, study)
    html_text = build_html(analysis, lang, figures, img_src=lambda k: f"fig://{k}",
                           study=study)
    d = docx.Document()
    from docx.shared import RGBColor
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].font.color.rgb = RGBColor(0x14, 0x14, 0x13)
    # the interface's look: serif headings, the title in ink, sections in terracotta
    for name, colour in (("Title", (0x14, 0x14, 0x13)), ("Heading 1", (0xC6, 0x61, 0x3F)),
                         ("Heading 2", (0x14, 0x14, 0x13))):
        style = d.styles[name]
        style.font.name = "Georgia"
        style.font.bold = False
        style.font.color.rgb = RGBColor(*colour)

    tokens = re.split(r"(<h1>.*?</h1>|<h2[^>]*>.*?</h2>|<h3>.*?</h3>|<table[^>]*>.*?</table>|"
                      r"<p[^>]*>.*?</p>|<li>.*?</li>)", html_text, flags=re.S)

    def strip(s):
        return html.unescape(re.sub(r"<[^>]+>", "", re.sub(r"<br\s*/?>", "\n", s))).strip()

    for tok in tokens:
        if tok.startswith("<h1>"):
            d.add_heading(strip(tok), level=0)
        elif tok.startswith("<h2"):
            d.add_heading(strip(tok), level=1)
        elif tok.startswith("<h3>"):
            d.add_heading(strip(tok), level=2)
        elif tok.startswith("<li>"):
            d.add_paragraph(strip(tok), style="List Bullet")
        elif tok.startswith("<table"):
            rows = re.findall(r"<tr>(.*?)</tr>", tok, flags=re.S)
            cells = [re.findall(r"<t[hd][^>]*>(.*?)</t[hd]>", r, flags=re.S) for r in rows]
            if cells:
                table = d.add_table(rows=len(cells), cols=len(cells[0]))
                table.style = "Light Grid Accent 1"
                for i, row in enumerate(cells):
                    for j, c in enumerate(row):
                        if j < len(table.columns):
                            table.cell(i, j).text = strip(c)
        elif tok.startswith("<p"):
            m = re.search(r"src='fig://([a-z_]+)'", tok)
            if m:
                d.add_picture(io.BytesIO(figures[m.group(1)]), width=Inches(6.3))
            else:
                text = strip(tok)
                if text:
                    d.add_paragraph(text)
    d.save(path)


def export_report(path: str, analysis, lang: str, study=None) -> None:
    ext = path.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        export_pdf(path, analysis, lang, study)
    elif ext == "docx":
        export_docx(path, analysis, lang, study)
    else:
        export_html(path if ext in ("html", "htm") else path + ".html", analysis, lang, study)
